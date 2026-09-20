"""Script to generate products.xlsx golden fixture."""
import os
import openpyxl
from openpyxl import Workbook

wb = Workbook()

# Products sheet
ws1 = wb.active
ws1.title = "Products"
ws1.append(["product_id", "product_name", "category", "subcategory", "unit_price", "launch_date"])
products = [
    ("P001", "DataPro Basic", "Software", "Analytics", 5000.00, "2022-06-01"),
    ("P002", "DataPro Enterprise", "Software", "Analytics", 12000.00, "2022-08-15"),
    ("P003", "CloudSync", "Software", "Integration", 8500.00, "2023-01-10"),
    ("P004", "MobileHub", "Mobile", "App", 3200.00, "2023-03-20"),
    ("P005", "SecureVault", "Security", "Compliance", 7500.00, "2023-07-05"),
    ("P006", "AI Insights", "AI/ML", "Analytics", 22000.00, "2024-01-15"),
]
for row in products:
    ws1.append(list(row))

# Pricing Tiers sheet (second usable sheet)
ws2 = wb.create_sheet("Pricing Tiers")
ws2.append(["tier", "discount_pct", "min_quantity", "notes"])
ws2.append(["Bronze", 0, 1, "Standard pricing"])
ws2.append(["Silver", 5, 3, "5% volume discount"])
ws2.append(["Gold", 10, 5, "10% loyalty discount"])

# Notes sheet (non-tabular, should be skipped)
ws3 = wb.create_sheet("Notes")
ws3["A1"] = "Internal product catalog - last updated 2025-01"
ws3["A2"] = "Do not distribute externally"

out_path = os.path.join(os.path.dirname(__file__), "products.xlsx")
wb.save(out_path)
print(f"Written: {out_path}")
